from pathlib import Path

import docx
import pytest
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Emu
from fastapi.testclient import TestClient

from app.converters.pdf_to_docx import convert_pdf
from app.converters.pptx_to_docx import convert_pptx
from app.main import app
from app.report import PasswordError
from tests.make_fixtures import FIX, make_pdf, make_pptx


@pytest.fixture(scope="session", autouse=True)
def fixtures():
    FIX.mkdir(exist_ok=True)
    if not (FIX / "sample.pdf").exists():
        make_pdf(FIX / "sample.pdf")
    if not (FIX / "locked.pdf").exists():
        make_pdf(FIX / "locked.pdf", password="secret")
    if not (FIX / "sample.pptx").exists():
        make_pptx(FIX / "sample.pptx")
    (FIX / "not-a-doc.txt").write_text("hello")


def _all_tables(d):
    """Top-level tables plus tables nested in cells (side-by-side slide content uses a layout table)."""
    found = []
    def walk(tables):
        for t in tables:
            found.append(t)
            for row in t.rows:
                for c in row.cells:
                    walk(c.tables)
    walk(d.tables)
    return found


def _text(d):
    cells = "\n".join(c.text for t in d.tables for row in t.rows for c in row.cells)
    return "\n".join(p.text for p in d.paragraphs) + "\n" + cells


def test_pdf(tmp_path: Path):
    out = tmp_path / "out.docx"
    report = convert_pdf(FIX / "sample.pdf", out)
    d = docx.Document(str(out))
    assert report.pages == 2
    assert "Quarterly report" in _text(d)
    assert "Lorem ipsum" in _text(d)
    assert len(d.tables) >= 1
    assert len(d.inline_shapes) >= 1
    assert any("no text layer" in w for w in report.warnings)
    fills = set()
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                shd = c._tc.tcPr.find(qn("w:shd")) if c._tc.tcPr is not None else None
                if shd is not None and shd.get(qn("w:fill")):
                    fills.add(shd.get(qn("w:fill")).upper())
    assert fills, "expected the 5% black header fill to survive as shading"
    for fill in fills:  # every shaded cell must be light, never solid black
        r, g, b = (int(fill[i:i + 2], 16) for i in (0, 2, 4))
        assert min(r, g, b) > 200, fill


def test_pdf_password(tmp_path: Path):
    with pytest.raises(PasswordError):
        convert_pdf(FIX / "locked.pdf", tmp_path / "x.docx")
    with pytest.raises(PasswordError):
        convert_pdf(FIX / "locked.pdf", tmp_path / "x.docx", password="wrong")
    report = convert_pdf(FIX / "locked.pdf", tmp_path / "ok.docx", password="secret")
    assert report.pages == 2


def test_pptx(tmp_path: Path):
    out = tmp_path / "out.docx"
    report = convert_pptx(FIX / "sample.pptx", out, include_notes=True)
    d = docx.Document(str(out))
    assert report.pages == 4
    heads = [p.text for p in d.paragraphs if p.style.name == "Heading 1"]
    assert heads == ["Agenda", "Numbers", "Slide 3", "Lists"]
    by_text = {p.text: p.style.name for p in d.paragraphs}
    assert by_text["No bullet here"] == "Normal"
    assert by_text["Bulleted box"] == "List Bullet"
    assert by_text["Numbered"] == "List Number"
    styles = {p.style.name for p in d.paragraphs}
    assert {"List Bullet", "List Bullet 2", "List Bullet 3"} <= styles
    bold = [r for p in d.paragraphs for r in p.runs if r.text == "Bold red"][0]
    assert bold.bold and str(bold.font.color.rgb) == "C00000" and bold.font.size.pt == 20
    # Inherited formatting: the title size, font and colour come from the master, not the run.
    title_run = [p for p in d.paragraphs if p.text == "Agenda"][0].runs[0]
    assert title_run.font.size is not None and title_run.font.size.pt >= 36
    assert title_run.font.name and title_run.font.color.rgb is not None
    first_point = [p for p in d.paragraphs if p.text == "First point"][0].runs[0]
    assert first_point.font.size is not None and first_point.font.size.pt >= 18
    assert first_point.font.name
    numbered = [p for p in d.paragraphs if p.text == "Numbered"][0]
    assert numbered.alignment == WD_ALIGN_PARAGRAPH.CENTER
    # A filled text box keeps its background, so its white text stays readable.
    bulleted = [p for p in d.paragraphs if p.text == "Bulleted box"][0]
    shd = bulleted._p.pPr.find(qn("w:shd"))
    assert shd is not None and shd.get(qn("w:fill")) == "1F5C8A"
    assert str(bulleted.runs[0].font.color.rgb) == "FFFFFF"
    grids = [t for t in _all_tables(d) if t.style is not None and t.style.name == "Table Grid"]
    assert len(grids) == 1 and report.tables == 1
    t = grids[0]
    assert t.cell(0, 0)._tc is t.cell(0, 1)._tc  # merged
    assert t.cell(2, 2).text == "r2c2"
    assert len(d.inline_shapes) == 2 and report.images == 2
    upright, rotated = d.inline_shapes
    assert upright.width > upright.height            # 120x80 source as is
    assert rotated.height > rotated.width            # same source turned by 90 degrees
    assert "Text inside a group" in _text(d)
    assert "Speaker notes" in [p.text for p in d.paragraphs if p.style.name == "Heading 3"]
    assert "Remember to thank the team." in _text(d)
    sec = d.sections[0]
    assert sec.page_width == Emu(9144000) and sec.page_height == Emu(6858000)


def test_pptx_without_notes(tmp_path: Path):
    out = tmp_path / "out.docx"
    convert_pptx(FIX / "sample.pptx", out, include_notes=False)
    d = docx.Document(str(out))
    assert "Remember to thank the team." not in _text(d)


def test_api_roundtrip():
    client = TestClient(app)
    assert client.get("/").status_code == 200
    with open(FIX / "sample.pptx", "rb") as f:
        r = client.post("/api/convert", files={"file": ("deck.pptx", f)}, data={"notes": "true"})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["filename"] == "deck.docx" and body["report"]["pages"] == 4
    dl = client.get(f"/api/result/{body['id']}")
    assert dl.status_code == 200 and dl.content[:2] == b"PK"
    assert "deck.docx" in dl.headers["content-disposition"]


def test_api_errors():
    client = TestClient(app)
    with open(FIX / "not-a-doc.txt", "rb") as f:
        assert client.post("/api/convert", files={"file": ("x.txt", f)}).status_code == 415
    with open(FIX / "locked.pdf", "rb") as f:
        r = client.post("/api/convert", files={"file": ("locked.pdf", f)})
    assert r.status_code == 401 and r.json()["error"] == "password"
    with open(FIX / "locked.pdf", "rb") as f:
        r = client.post("/api/convert", files={"file": ("locked.pdf", f)}, data={"password": "secret"})
    assert r.status_code == 200
    assert client.get("/api/result/doesnotexist").status_code == 404

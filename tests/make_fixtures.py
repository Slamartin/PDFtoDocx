"""Generates the sample inputs used by the tests. Run: uv run python tests/make_fixtures.py"""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

import pymupdf
from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

FIX = Path(__file__).parent / "fixtures"


def png_bytes() -> bytes:
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 120, 80), False)
    pix.clear_with(90)
    return pix.tobytes("png")


def make_pdf(path: Path, password: str | None = None) -> None:
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Quarterly report", fontsize=22, fontname="helv")
    lorem = ("Lorem ipsum dolor sit amet, consectetur adipiscing elit. Sed do eiusmod tempor "
             "incididunt ut labore et dolore magna aliqua. Ut enim ad minim veniam. ") * 3
    page.insert_textbox(pymupdf.Rect(72, 100, 290, 300), lorem, fontsize=10, fontname="helv")
    page.insert_textbox(pymupdf.Rect(310, 100, 530, 300), lorem, fontsize=10, fontname="tiro")
    y0 = 320
    # Header row drawn like web-to-PDF output: black at 5% opacity, which should become light grey, not black.
    page.draw_rect(pymupdf.Rect(72, y0, 72 + 360, y0 + 24), fill=(0, 0, 0), fill_opacity=0.05, width=0)
    for r in range(3):
        for c in range(3):
            rect = pymupdf.Rect(72 + c * 120, y0 + r * 24, 72 + (c + 1) * 120, y0 + (r + 1) * 24)
            page.draw_rect(rect, color=(0, 0, 0), width=0.8)
            page.insert_text((rect.x0 + 4, rect.y1 - 7), f"R{r + 1}C{c + 1}", fontsize=10, fontname="helv")
    page.insert_image(pymupdf.Rect(72, 420, 252, 540), stream=png_bytes())
    doc.new_page()  # blank page: no text layer
    kwargs = {}
    if password:
        kwargs = dict(encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw=password, owner_pw=password)
    doc.save(str(path), **kwargs)
    doc.close()


def make_pptx(path: Path) -> None:
    prs = Presentation()
    # Slide 1: title + bullets with levels and run formatting
    s1 = prs.slides.add_slide(prs.slide_layouts[1])
    s1.shapes.title.text = "Agenda"
    tf = s1.placeholders[1].text_frame
    tf.text = "First point"
    p = tf.add_paragraph(); p.text = "Nested point"; p.level = 1
    p = tf.add_paragraph(); p.text = "Deeper point"; p.level = 2
    p = tf.add_paragraph()
    r = p.add_run(); r.text = "Bold red"; r.font.bold = True; r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00); r.font.size = Pt(20)
    # Slide 2: table with a merged cell + picture
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "Numbers"
    tbl = s2.shapes.add_table(3, 3, Inches(0.5), Inches(1.5), Inches(5), Inches(1.5)).table
    for r_ in range(3):
        for c in range(3):
            tbl.cell(r_, c).text = f"r{r_}c{c}"
    tbl.cell(0, 0).merge(tbl.cell(0, 1))
    s2.shapes.add_picture(BytesIO(png_bytes()), Inches(6), Inches(1.5), width=Inches(3))
    rotated = s2.shapes.add_picture(BytesIO(png_bytes()), Inches(6), Inches(4), width=Inches(3))
    rotated.rotation = 90.0
    # Slide 3: no title, group shape, speaker notes
    s3 = prs.slides.add_slide(prs.slide_layouts[6])
    grp = s3.shapes.add_group_shape()
    tb = grp.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    tb.text_frame.text = "Text inside a group"
    s3.notes_slide.notes_text_frame.text = "Remember to thank the team."
    # Slide 4: bullets decided by list styles rather than paragraph properties
    s4 = prs.slides.add_slide(prs.slide_layouts[1])
    s4.shapes.title.text = "Lists"
    body = s4.placeholders[1]
    body.text_frame.text = "No bullet here"
    lst = body.text_frame._txBody.find(qn("a:lstStyle"))
    if lst is None:
        lst = etree.SubElement(body.text_frame._txBody, qn("a:lstStyle"))
    etree.SubElement(etree.SubElement(lst, qn("a:lvl1pPr")), qn("a:buNone"))
    box = s4.shapes.add_textbox(Inches(1), Inches(3), Inches(5), Inches(1))
    box.text_frame.text = "Bulleted box"
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(0x1F, 0x5C, 0x8A)
    box.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    lst = box.text_frame._txBody.find(qn("a:lstStyle"))
    lvl = etree.SubElement(lst, qn("a:lvl1pPr"))
    etree.SubElement(lvl, qn("a:buChar")).set("char", "\u2022")
    p = box.text_frame.add_paragraph()
    p.text = "Numbered"
    p.alignment = PP_ALIGN.CENTER
    etree.SubElement(p._p.get_or_add_pPr(), qn("a:buAutoNum")).set("type", "arabicPeriod")
    prs.save(str(path))


if __name__ == "__main__":
    FIX.mkdir(exist_ok=True)
    make_pdf(FIX / "sample.pdf")
    make_pdf(FIX / "locked.pdf", password="secret")
    make_pptx(FIX / "sample.pptx")
    (FIX / "not-a-doc.txt").write_text("hello")
    print("fixtures written to", FIX)

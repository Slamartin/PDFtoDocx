"""PPTX -> DOCX: one slide per page, content rebuilt as editable Word elements.

Formatting in PowerPoint is mostly inherited: a run with no size gets it from the shape's
list style, the layout placeholder, the master placeholder, the master text styles or the
presentation defaults, and colours/fonts may be theme references. The resolver below walks
that chain so the Word output carries the sizes, fonts, colours and alignment a viewer sees.
"""

from __future__ import annotations

import colorsys
import logging
import math
from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.image.exceptions import UnrecognizedImageError
from docx.shared import Cm, Emu, Pt, RGBColor
from lxml import etree
from PIL import Image, ImageOps
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.exc import PackageNotFoundError
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.shapes.graphfrm import GraphicFrame
from pptx.shapes.picture import Picture
from pptx.text.text import _Run

from app.report import ConversionReport, CorruptFileError

log = logging.getLogger(__name__)

EMU_PER_INCH = 914400
MIN_PAGE = int(0.5 * EMU_PER_INCH)
MAX_PAGE = 22 * EMU_PER_INCH
MARGIN = Cm(1.27)
ROW_GRID = Cm(0.5)
MAX_PICTURE_HEIGHT = 0.85  # share of the usable page height; the per-slide budget shrinks pictures further when needed
MAX_GAP = 0.2  # one vertical gap between shapes is at most this share of the usable height
PAGE_FILL = 0.92  # share of the usable height a slide's content may fill before pictures and gaps shrink
CHAR_WIDTH = 0.6  # average glyph width relative to the font size, for wrapping estimates
HEADING_FALLBACK_HEIGHT = 450000  # EMU reserved for the 'Slide n' heading when a slide has no title
EMU_PER_PT = 12700
LINE_FACTOR = 1.25  # Word line height relative to the font size, used for height estimates

TITLE_TYPES = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.VERTICAL_TITLE}
CHROME_TYPES = {PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER}
RASTER_EXT = {"png", "jpg", "jpeg", "gif", "bmp", "tif", "tiff"}

ALIGN = {
    "l": WD_ALIGN_PARAGRAPH.LEFT,
    "ctr": WD_ALIGN_PARAGRAPH.CENTER,
    "r": WD_ALIGN_PARAGRAPH.RIGHT,
    "just": WD_ALIGN_PARAGRAPH.JUSTIFY,
    "dist": WD_ALIGN_PARAGRAPH.JUSTIFY,
}
PRESET_COLORS = {
    "black": "000000",
    "white": "FFFFFF",
    "red": "FF0000",
    "blue": "0000FF",
    "green": "008000",
    "gray": "808080",
    "grey": "808080",
}


def _pct(raw: str | None) -> float | None:
    """OOXML percentages come as thousandths of a percent ('150000') or as a string ('150%'); both -> 1.5.
    None when the attribute is missing or malformed."""
    if not raw:
        return None
    raw = raw.strip()
    try:
        if raw.endswith("%"):
            return float(raw[:-1]) / 100
        return int(raw) / 100000
    except ValueError:
        return None


# ---------- theme and inherited style resolution ----------


class Theme:
    """Colour scheme, font scheme, master text styles and presentation defaults for one slide master."""

    def __init__(self, prs, master):
        self.colors: dict[str, str] = {}
        self.major_font = None
        self.minor_font = None
        self.clr_map: dict[str, str] = {}
        self.tx_styles: dict[str, object] = {}
        self.default_text_style = None
        try:
            root = etree.fromstring(master.part.part_related_by(RT.THEME).blob)
            scheme = root.find(".//" + qn("a:clrScheme"))
            if scheme is not None:
                for child in scheme:
                    name = etree.QName(child).localname
                    if len(child):
                        self.colors[name] = self._base_color(child[0])
            fonts = root.find(".//" + qn("a:fontScheme"))
            if fonts is not None:
                major = fonts.find(qn("a:majorFont") + "/" + qn("a:latin"))
                minor = fonts.find(qn("a:minorFont") + "/" + qn("a:latin"))
                self.major_font = major.get("typeface") if major is not None else None
                self.minor_font = minor.get("typeface") if minor is not None else None
        except Exception:  # a deck without a readable theme still converts, just without theme colours
            log.debug("theme not readable", exc_info=True)
        clr_map = master._element.find(qn("p:clrMap"))
        if clr_map is not None:
            self.clr_map = dict(clr_map.attrib)
        tx = master._element.find(qn("p:txStyles"))
        if tx is not None:
            for name in ("titleStyle", "bodyStyle", "otherStyle"):
                self.tx_styles[name] = tx.find(qn("p:" + name))
        self.default_text_style = prs._element.find(qn("p:defaultTextStyle"))

    @staticmethod
    def _base_color(el):
        tag = etree.QName(el).localname
        if tag == "srgbClr":
            return el.get("val")
        if tag == "sysClr":
            return el.get("lastClr")
        if tag == "prstClr":
            return PRESET_COLORS.get(el.get("val"))
        return None

    def color(self, fill_parent):
        """Resolves the solid fill of an rPr/defRPr/spPr element to 'RRGGBB', or None."""
        solid = fill_parent.find(qn("a:solidFill"))
        if solid is None or not len(solid):
            return None
        return self.color_el(solid[0])

    def color_el(self, el):
        """Resolves one colour element (srgbClr, schemeClr, sysClr, prstClr) to 'RRGGBB', or None."""
        tag = etree.QName(el).localname
        if tag == "schemeClr":
            name = el.get("val")
            base = self.colors.get(self.clr_map.get(name, name))
        else:
            base = self._base_color(el)
        if not base or len(base) != 6:
            return None
        return self._modify(base, el)

    @staticmethod
    def _modify(hex_color: str, el) -> str:
        r, g, b = (int(hex_color[i : i + 2], 16) / 255 for i in (0, 2, 4))
        for mod in el:
            name = etree.QName(mod).localname
            val = _pct(mod.get("val", "100000"))
            if val is None:
                continue
            if name == "lumMod" or name == "lumOff":
                h, lum, s = colorsys.rgb_to_hls(r, g, b)
                lum = lum * val if name == "lumMod" else lum + val
                r, g, b = colorsys.hls_to_rgb(h, min(1.0, max(0.0, lum)), s)
            elif name == "tint":
                r, g, b = (1 - (1 - c) * val for c in (r, g, b))
            elif name == "shade":
                r, g, b = (c * val for c in (r, g, b))
        return "".join(f"{int(round(min(1.0, max(0.0, c)) * 255)):02X}" for c in (r, g, b))

    def font(self, typeface: str | None):
        if not typeface:
            return None
        if typeface.startswith("+mj"):
            return self.major_font
        if typeface.startswith("+mn"):
            return self.minor_font
        return typeface

    def text_color_default(self):
        return self.colors.get(self.clr_map.get("tx1", "dk1"))


@dataclass
class RunStyle:
    size: float | None = None
    bold: bool | None = None
    italic: bool | None = None
    underline: bool | None = None
    color: str | None = None
    font: str | None = None


@dataclass
class ParaStyle:
    align: str | None = None
    line_pct: float | None = None
    line_pts: float | None = None
    before: float | None = None
    after: float | None = None
    before_pct: float | None = None
    after_pct: float | None = None
    mar_left: int | None = None
    indent: int | None = None
    bullet: str = "plain"  # 'bullet', 'number' or 'plain'


@dataclass
class Ctx:
    doc: Document
    report: ConversionReport
    usable_width: int
    usable_height: int
    prs: object = None
    slide_no: int = 0
    slide: object = None
    theme: Theme | None = None
    themes: dict = field(default_factory=dict)
    font_ref_color: str | None = None  # text colour from the current shape's style reference, if any
    cursor: int = 0  # vertical position on the current slide already covered by emitted content, EMU
    gap_budget: int = 0  # remaining vertical whitespace that may still be reproduced on this slide


def _placeholder_type(shape):
    try:
        return shape.placeholder_format.type if shape.is_placeholder else None
    except Exception:
        return None


def _lvl(style_parent, level: int):
    return None if style_parent is None else style_parent.find(qn(f"a:lvl{level + 1}pPr"))


def _lst_style(shape):
    try:
        if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
            return shape.text_frame._txBody.find(qn("a:lstStyle"))
    except Exception:
        pass
    return None


def _pPr_chain(paragraph, shape, ctx: Ctx) -> list:
    """Paragraph property elements in precedence order, most specific first."""
    level = paragraph.level or 0
    chain = [paragraph._p.pPr, _lvl(_lst_style(shape), level)]
    ptype = _placeholder_type(shape) if shape is not None else None
    if ptype is not None:
        base = shape
        for _ in range(2):  # layout placeholder, then master placeholder
            try:
                base = base._base_placeholder
            except Exception:
                base = None
            if base is None:
                break
            chain.append(_lvl(_lst_style(base), level))
    theme = ctx.theme
    if theme is not None:
        name = "titleStyle" if ptype in TITLE_TYPES else ("bodyStyle" if ptype is not None else "otherStyle")
        chain.append(_lvl(theme.tx_styles.get(name), level))
        chain.append(_lvl(theme.default_text_style, level))
    return [c for c in chain if c is not None]


def _bullet_from(pPr):
    if pPr.find(qn("a:buNone")) is not None:
        return "plain"
    if pPr.find(qn("a:buAutoNum")) is not None:
        return "number"
    if pPr.find(qn("a:buChar")) is not None or pPr.find(qn("a:buBlip")) is not None:
        return "bullet"
    return None


def _spacing_pts(el):
    """a:spcBef / a:spcAft / a:lnSpc -> (percent, points)."""
    pct = el.find(qn("a:spcPct"))
    if pct is not None:
        value = _pct(pct.get("val"))
        if value is not None:
            return value, None
    pts = el.find(qn("a:spcPts"))
    if pts is not None and pts.get("val"):
        try:
            return None, int(pts.get("val")) / 100
        except ValueError:
            pass
    return None, None


def _para_style(chain) -> ParaStyle:
    ps = ParaStyle()
    seen = set()
    for pPr in chain:
        if "align" not in seen and pPr.get("algn"):
            ps.align = pPr.get("algn")
            seen.add("align")
        if "marL" not in seen and pPr.get("marL") is not None:
            ps.mar_left = int(pPr.get("marL"))
            seen.add("marL")
        if "indent" not in seen and pPr.get("indent") is not None:
            ps.indent = int(pPr.get("indent"))
            seen.add("indent")
        if "ln" not in seen:
            ln = pPr.find(qn("a:lnSpc"))
            if ln is not None:
                ps.line_pct, ps.line_pts = _spacing_pts(ln)
                seen.add("ln")
        for key, tag in (("before", "a:spcBef"), ("after", "a:spcAft")):
            if key not in seen:
                el = pPr.find(qn(tag))
                if el is not None:
                    pct, pts = _spacing_pts(el)
                    setattr(ps, key, pts or 0.0)
                    setattr(ps, key + "_pct", pct)
                    seen.add(key)
        if "bullet" not in seen:
            kind = _bullet_from(pPr)
            if kind:
                ps.bullet = kind
                seen.add("bullet")
    return ps


def _run_style(rPr, chain, ctx: Ctx, font_scale: float = 1.0) -> RunStyle:
    elems = [rPr] if rPr is not None else []
    elems += [d for d in (pPr.find(qn("a:defRPr")) for pPr in chain) if d is not None]
    rs = RunStyle()
    for el in elems:
        if rs.size is None and el.get("sz"):
            rs.size = int(el.get("sz")) / 100 * font_scale
        if rs.bold is None and el.get("b") is not None:
            rs.bold = el.get("b") in ("1", "true")
        if rs.italic is None and el.get("i") is not None:
            rs.italic = el.get("i") in ("1", "true")
        if rs.underline is None and el.get("u") is not None:
            rs.underline = el.get("u") != "none"
        if (
            rs.color is None
            and ctx.theme is not None
            and (el.find(qn("a:solidFill")) is not None or el.find(qn("a:noFill")) is not None)
        ):
            rs.color = ctx.theme.color(el) or rs.color
            if rs.color is None:
                rs.color = ""  # explicit fill that could not be resolved: stop looking further up
        if rs.font is None:
            latin = el.find(qn("a:latin"))
            if latin is not None and latin.get("typeface"):
                rs.font = ctx.theme.font(latin.get("typeface")) if ctx.theme else latin.get("typeface")
    if not rs.color and ctx.font_ref_color:
        rs.color = ctx.font_ref_color
    if ctx.theme is not None:
        if not rs.color:
            rs.color = ctx.theme.text_color_default()
        if rs.font is None:
            rs.font = ctx.theme.minor_font
    return rs


def _shape_style_color(shape, ctx: Ctx, ref_tag: str) -> str | None:
    """Colour of a p:style reference (a:fillRef or a:fontRef) on the shape, if any."""
    if ctx.theme is None:
        return None
    try:
        style = shape._element.find(qn("p:style"))
        ref = style.find(qn(ref_tag)) if style is not None else None
        if ref is None or not len(ref):
            return None
        if ref_tag == "a:fillRef" and int(ref.get("idx", "0") or 0) == 0:
            return None  # idx 0 means no fill
        return ctx.theme.color_el(ref[0])
    except Exception:
        return None


def _shape_fill(shape, ctx: Ctx) -> str | None:
    """Solid background of a shape: its own spPr, inherited placeholder spPr, or its style's fill reference."""
    if ctx.theme is None:
        return None
    candidates = [shape]
    base = shape
    for _ in range(2):
        try:
            base = base._base_placeholder
        except Exception:
            base = None
        if base is None:
            break
        candidates.append(base)
    for cand in candidates:
        try:
            spPr = cand._element.find(qn("p:spPr"))
        except Exception:
            spPr = None
        if spPr is None:
            continue
        if spPr.find(qn("a:noFill")) is not None:
            return None
        color = ctx.theme.color(spPr)
        if color:
            return color
    return _shape_style_color(shape, ctx, "a:fillRef")


def _insert_in_order(parent, new, successors) -> None:
    """Inserts `new` before the first child whose tag is in `successors`, else appends (OOXML order matters)."""
    tags = {qn(t) for t in successors}
    for child in parent:
        if child.tag in tags:
            child.addprevious(new)
            return
    parent.append(new)


PPR_AFTER_SHD = (
    "w:tabs",
    "w:suppressAutoHyphens",
    "w:kinsoku",
    "w:wordWrap",
    "w:overflowPunct",
    "w:topLinePunct",
    "w:autoSpaceDE",
    "w:autoSpaceDN",
    "w:bidi",
    "w:adjustRightInd",
    "w:snapToGrid",
    "w:spacing",
    "w:ind",
    "w:contextualSpacing",
    "w:mirrorIndents",
    "w:suppressOverlap",
    "w:jc",
    "w:textDirection",
    "w:textAlignment",
    "w:textboxTightWrap",
    "w:outlineLvl",
    "w:divId",
    "w:cnfStyle",
    "w:rPr",
    "w:sectPr",
    "w:pPrChange",
)
TCPR_AFTER_SHD = (
    "w:noWrap",
    "w:tcMar",
    "w:textDirection",
    "w:tcFitText",
    "w:vAlign",
    "w:hideMark",
    "w:headers",
    "w:cellIns",
    "w:cellDel",
    "w:cellMerge",
    "w:tcPrChange",
)


def _shade_paragraph(p, color: str) -> None:
    pPr = p._p.get_or_add_pPr()
    old = pPr.find(qn("w:shd"))
    if old is not None:
        pPr.remove(old)
    shd = etree.Element(qn("w:shd"))
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    _insert_in_order(pPr, shd, PPR_AFTER_SHD)


def _shade_cell(cell, color: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    old = tcPr.find(qn("w:shd"))
    if old is not None:
        tcPr.remove(old)
    shd = etree.Element(qn("w:shd"))
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    _insert_in_order(tcPr, shd, TCPR_AFTER_SHD)


def _is_symbol_char(ch: str) -> bool:
    return "\uf000" <= ch <= "\uf0ff"


def _split_symbol_runs(text: str, sym_font: str | None):
    """Yields (text, font_override) pieces so that symbol-font characters get that font and the rest keeps its own."""
    if not sym_font or not any(_is_symbol_char(c) for c in text):
        yield text, None
        return
    piece, piece_sym = "", _is_symbol_char(text[0])
    for ch in text:
        is_sym = _is_symbol_char(ch) or (ch == " " and piece_sym)
        if is_sym != piece_sym and piece:
            yield piece, sym_font if piece_sym else None
            piece = ""
        piece += ch
        piece_sym = is_sym
    if piece:
        yield piece, sym_font if piece_sym else None


def _apply_run_style(run_docx, rs: RunStyle) -> None:
    run_docx.bold = bool(rs.bold)
    run_docx.italic = bool(rs.italic)
    if rs.underline is not None:
        run_docx.underline = rs.underline
    if rs.size:
        run_docx.font.size = Pt(max(1.0, rs.size))
    if rs.font:
        run_docx.font.name = rs.font
    if rs.color and len(rs.color) == 6:
        run_docx.font.color.rgb = RGBColor.from_string(rs.color)


def _font_scale(shape) -> float:
    """normAutofit shrinks text on the slide; PowerPoint stores the applied scale."""
    try:
        body_pr = shape.text_frame._txBody.find(qn("a:bodyPr"))
        auto = body_pr.find(qn("a:normAutofit")) if body_pr is not None else None
        scale = _pct(auto.get("fontScale")) if auto is not None else None
        if scale:
            return scale
    except Exception:
        pass
    return 1.0


# ---------- page setup ----------


def _setup_page(doc, prs, report: ConversionReport) -> tuple[int, int]:
    w, h = int(prs.slide_width), int(prs.slide_height)
    cw, ch = min(MAX_PAGE, max(MIN_PAGE, w)), min(MAX_PAGE, max(MIN_PAGE, h))
    if (cw, ch) != (w, h):
        report.warn("Slide size is outside Word's page limits and was clamped.")
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE if cw >= ch else WD_ORIENT.PORTRAIT
    sec.page_width = Emu(cw)
    sec.page_height = Emu(ch)
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = MARGIN
    return cw - 2 * MARGIN, ch - 2 * MARGIN


# ---------- reading order ----------


def _child_space(group) -> tuple[int, int, float, float]:
    """How a group's children map onto its own box: (chOff x, chOff y, scale x, scale y).

    A child at (x, y) in the group's child space sits at group.off + (x - chOff) * scale in the
    parent's space, where scale = ext / chExt. PowerPoint sets chOff = off and chExt = ext when a
    group is made, so the child coordinates usually are slide coordinates already; moving or
    resizing the group changes only off/ext and the children stay where they were."""
    try:
        xfrm = group._element.find(qn("p:grpSpPr")).find(qn("a:xfrm"))
        ch_off, ch_ext, ext = xfrm.find(qn("a:chOff")), xfrm.find(qn("a:chExt")), xfrm.find(qn("a:ext"))
        cx, cy = int(ch_ext.get("cx")), int(ch_ext.get("cy"))
        sx = int(ext.get("cx")) / cx if cx else 1.0
        sy = int(ext.get("cy")) / cy if cy else 1.0
        return int(ch_off.get("x")), int(ch_off.get("y")), sx, sy
    except Exception:
        return 0, 0, 1.0, 1.0


def _flatten(shapes, acc: list, base_top: int, base_left: int, scale_y: float = 1.0, scale_x: float = 1.0) -> None:
    """Collects (top, left, shape) in slide coordinates, descending into groups.
    Shape sizes are left in their own (child) space; a scaled group is rare and only shifts the estimates."""
    for shape in shapes:
        top = base_top + int((shape.top or 0) * scale_y)
        left = base_left + int((shape.left or 0) * scale_x)
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            ch_x, ch_y, sx, sy = _child_space(shape)
            inner_y, inner_x = scale_y * sy, scale_x * sx
            _flatten(shape.shapes, acc, top - int(ch_y * inner_y), left - int(ch_x * inner_x), inner_y, inner_x)
        else:
            acc.append((top, left, shape))


def _ordered_shapes(slide):
    acc: list = []
    _flatten(slide.shapes, acc, 0, 0)
    title = slide.shapes.title
    body = [t for t in acc if t[2] is not title and _placeholder_type(t[2]) not in TITLE_TYPES]
    titles = [t for t in acc if t not in body]
    body.sort(key=lambda t: (round(t[0] / ROW_GRID), t[1]))
    return titles, body


def _take_gap(ctx: Ctx, top: int) -> int:
    """Vertical whitespace between the content emitted so far and this shape, in EMU."""
    gap = top - ctx.cursor
    if gap <= 0:
        return 0
    gap = min(gap, int(ctx.usable_height * MAX_GAP), ctx.gap_budget)
    ctx.gap_budget -= gap
    return gap


def _advance(ctx: Ctx, top: int, height: int) -> None:
    ctx.cursor = max(ctx.cursor, top + (height or 0))


# ---------- text ----------


def _paragraph_is_empty(paragraph) -> bool:
    return not paragraph.text.strip() and not any(etree.QName(c).localname == "br" for c in paragraph._p)


def _trimmed_paragraphs(text_frame):
    paras = list(text_frame.paragraphs)
    while paras and _paragraph_is_empty(paras[0]):
        paras.pop(0)
    while paras and _paragraph_is_empty(paras[-1]):
        paras.pop()
    return paras


def _first_run_rPr(paragraph):
    for child in paragraph._p:
        if etree.QName(child).localname in ("r", "fld"):
            return child.find(qn("a:rPr"))
    return paragraph._p.find(qn("a:endParaRPr"))


def _set_paragraph_mark_size(p, size_pt: float) -> None:
    """The bullet or number glyph takes the paragraph mark's size, not the runs' size."""
    pPr = p._p.get_or_add_pPr()
    rPr = pPr.find(qn("w:rPr"))
    if rPr is None:
        rPr = etree.SubElement(pPr, qn("w:rPr"))
    for tag in ("w:sz", "w:szCs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set(qn("w:val"), str(int(round(size_pt * 2))))


def _emit_paragraph(
    container, paragraph, shape, ctx: Ctx, reuse=None, force_style=None, font_scale=1.0, extra_before=0, shade=None
):
    chain = _pPr_chain(paragraph, shape, ctx)
    ps = _para_style(chain)
    first_rs = _run_style(_first_run_rPr(paragraph), chain, ctx, font_scale)
    size = first_rs.size or 18.0
    level = min(paragraph.level or 0, 2)
    if force_style:
        style = force_style
    elif ps.bullet == "bullet":
        style = "List Bullet" if level == 0 else f"List Bullet {level + 1}"
    elif ps.bullet == "number":
        style = "List Number" if level == 0 else f"List Number {level + 1}"
    else:
        style = "Normal"
    if reuse is not None:
        p = reuse
        p.style = ctx.doc.styles[style]
    else:
        p = container.add_paragraph(style=style)

    pf = p.paragraph_format
    if ps.align in ALIGN:
        p.alignment = ALIGN[ps.align]
    pf.space_before = Pt((ps.before or 0) + (ps.before_pct or 0) * size) + Emu(extra_before)
    pf.space_after = Pt((ps.after or 0) + (ps.after_pct or 0) * size)
    if ps.bullet != "plain" and not force_style:
        _set_paragraph_mark_size(p, size)
    if shade:
        _shade_paragraph(p, shade)
    if ps.line_pct:
        pf.line_spacing = ps.line_pct
    elif ps.line_pts:
        pf.line_spacing = Pt(ps.line_pts)
    else:
        pf.line_spacing = 1.0
    if ps.bullet == "plain" and not force_style:
        if ps.mar_left:
            pf.left_indent = Emu(max(0, ps.mar_left))
        if ps.indent:
            pf.first_line_indent = Emu(ps.indent)

    if _paragraph_is_empty(paragraph):
        # Keep the blank line at the height PowerPoint gives it (endParaRPr carries the size).
        end = paragraph._p.find(qn("a:endParaRPr"))
        _apply_run_style(p.add_run(""), _run_style(end, chain, ctx, font_scale))
        return p

    for child in paragraph._p:
        tag = etree.QName(child).localname
        if tag == "r":
            run = _Run(child, paragraph)
            text = run.text
            try:
                address = run.hyperlink.address
            except Exception:
                address = None
            if address:
                text = f"{text} ({address})"
                ctx.report.note("Hyperlinks were flattened to text with the URL in brackets.")
            rPr = child.find(qn("a:rPr"))
            sym = rPr.find(qn("a:sym")) if rPr is not None else None
            rs = _run_style(rPr, chain, ctx, font_scale)
            for piece, font_override in _split_symbol_runs(text, sym.get("typeface") if sym is not None else None):
                r = p.add_run(piece)
                _apply_run_style(r, rs)
                if font_override:
                    r.font.name = font_override
        elif tag == "br":
            r = p.add_run()
            _apply_run_style(r, _run_style(child.find(qn("a:rPr")), chain, ctx, font_scale))
            r.add_break(WD_BREAK.LINE)
        elif tag == "fld":
            t = child.find(qn("a:t"))
            if t is not None and t.text:
                r = p.add_run(t.text)
                _apply_run_style(r, _run_style(child.find(qn("a:rPr")), chain, ctx, font_scale))
    return p


def _estimate_text_height(text_frame, shape, ctx: Ctx, width: int | None = None) -> int:
    """Rough height of a text frame once laid out in Word at the given width, in EMU."""
    scale = _font_scale(shape) if shape is not None else 1.0
    width = width or ctx.usable_width
    total = 0.0
    for paragraph in _trimmed_paragraphs(text_frame):
        chain = _pPr_chain(paragraph, shape, ctx)
        ps = _para_style(chain)
        rs = _run_style(_first_run_rPr(paragraph), chain, ctx, scale)
        size = rs.size or 18.0
        width_pt = max(50.0, (width - (ps.mar_left or 0)) / EMU_PER_PT)
        breaks = sum(1 for c in paragraph._p if etree.QName(c).localname == "br")
        lines = max(1, math.ceil((len(paragraph.text) + 1) * size * CHAR_WIDTH / width_pt)) + breaks
        line_h = ps.line_pts if ps.line_pts else size * (ps.line_pct or 1.0)
        total += lines * line_h * LINE_FACTOR
        total += (ps.before or 0) + (ps.before_pct or 0) * size + (ps.after or 0) + (ps.after_pct or 0) * size
    return int(total * EMU_PER_PT)


def _picture_box(shape, ctx: Ctx, max_width: int | None = None) -> tuple[int, int]:
    """Displayed (width, height) in EMU before the slide's height budget is applied."""
    w, h = int(shape.width or 0), int(shape.height or 0)
    if round(float(shape.rotation or 0)) % 180 == 90:
        w, h = h, w
    if w <= 0 or h <= 0:
        w, h = ctx.usable_width // 2, ctx.usable_height // 2
    width = min(w, max_width or ctx.usable_width)
    max_height = int(ctx.usable_height * MAX_PICTURE_HEIGHT)
    if h * width / w > max_height:
        width = int(max_height * w / h)
    return width, int(h * width / w)


# ---------- layout: boxes, rows and the per-slide height budget ----------


@dataclass
class Box:
    top: int
    left: int
    bottom: int
    right: int
    shape: object
    kind: str  # 'text', 'picture', 'table', 'other'
    col_width: int = 0  # width of the column this box is emitted into
    height: int = 0  # estimated height at full size, EMU
    pic_w: int = 0
    pic_h: int = 0


def _classify(shape) -> str | None:
    if _placeholder_type(shape) in CHROME_TYPES:
        return None
    if getattr(shape, "has_table", False) and shape.has_table:
        return "table"
    if getattr(shape, "has_chart", False) and shape.has_chart:
        return "other"
    if isinstance(shape, Picture):
        return "picture"
    if isinstance(shape, GraphicFrame) or shape.shape_type in (
        MSO_SHAPE_TYPE.EMBEDDED_OLE_OBJECT,
        MSO_SHAPE_TYPE.LINKED_OLE_OBJECT,
    ):
        return "other"
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame and shape.text_frame.text.strip():
        return "text"
    return None


def _text_extent_width(text_frame, shape, ctx: Ctx) -> int:
    """Width the longest line of a text frame needs, in EMU (0 when unknown)."""
    scale = _font_scale(shape) if shape is not None else 1.0
    widest = 0.0
    for paragraph in _trimmed_paragraphs(text_frame):
        chain = _pPr_chain(paragraph, shape, ctx)
        ps = _para_style(chain)
        size = _run_style(_first_run_rPr(paragraph), chain, ctx, scale).size or 18.0
        for line in paragraph.text.split("\x0b"):
            widest = max(widest, len(line) * size * CHAR_WIDTH * EMU_PER_PT + (ps.mar_left or 0))
    return int(widest)


def _make_boxes(body, ctx: Ctx) -> list[Box]:
    boxes = []
    for top, left, shape in body:
        kind = _classify(shape)
        if kind is None:
            if _placeholder_type(shape) in CHROME_TYPES:
                ctx.report.note("Slide numbers, dates and footers were left out.")
            continue
        w, h = max(int(shape.width or 0), 0), max(int(shape.height or 0), 0)
        if round(float(getattr(shape, "rotation", 0) or 0)) % 180 == 90:
            # A shape turned by 90 degrees occupies a box with swapped extents around the same centre.
            cx, cy = left + w // 2, top + h // 2
            left, top, w, h = cx - h // 2, cy - w // 2, h, w
        if kind == "text":
            try:
                extent = _text_extent_width(shape.text_frame, shape, ctx)
                if 0 < extent < w:
                    w = extent + Cm(0.5)
            except Exception:
                pass
        boxes.append(Box(top=top, left=left, bottom=top + h, right=left + w, shape=shape, kind=kind))
    return boxes


def _group_rows(boxes: list[Box]) -> list[list[Box]]:
    """Shapes that overlap vertically but not horizontally sit side by side on the slide;
    they become one row with a column per shape."""
    rows: list[list[Box]] = []
    for box in sorted(boxes, key=lambda b: (b.top, b.left)):
        placed = False
        for row in rows:
            r_top, r_bottom = min(b.top for b in row), max(b.bottom for b in row)
            overlap = min(r_bottom, box.bottom) - max(r_top, box.top)
            if overlap <= 0.3 * max(1, min(box.bottom - box.top, r_bottom - r_top)):
                continue
            side_by_side = all(
                min(b.right, box.right) - max(b.left, box.left)
                <= 0.25 * max(1, min(b.right - b.left, box.right - box.left))
                for b in row
            )
            if side_by_side and len(row) < 4:
                row.append(box)
                placed = True
                break
        if not placed:
            rows.append([box])
    for row in rows:
        row.sort(key=lambda b: b.left)
    rows.sort(key=lambda r: min(b.top for b in r))
    return rows


def _column_widths(row: list[Box], ctx: Ctx) -> list[int]:
    if len(row) == 1:
        return [ctx.usable_width]
    spans = []
    for i, box in enumerate(row):
        if i + 1 < len(row):
            spans.append(max(1, row[i + 1].left - box.left))
        else:
            spans.append(max(1, box.right - box.left))
    total = sum(spans)
    widths = [max(0.12, sp / total) for sp in spans]
    # A text column must hold its longest line, or the budget estimate explodes with wrapped lines.
    for i, box in enumerate(row):
        if box.kind == "text":
            try:
                need = (_text_extent_width(box.shape.text_frame, box.shape, ctx) + Cm(1.5)) / ctx.usable_width
            except Exception:
                need = 0
            widths[i] = max(widths[i], min(0.65, need))
    norm = sum(widths)
    return [int(ctx.usable_width * w / norm) for w in widths]


def _plan_slide(title_boxes: list[Box], rows: list[list[Box]], ctx: Ctx) -> None:
    """Estimates how much of the page the slide needs. Pictures and the whitespace between
    rows shrink so that one slide stays on one page where possible."""
    for row in [title_boxes] + rows:
        widths = _column_widths(row, ctx) if row else []
        for box, width in zip(row, widths):
            box.col_width = width
            if box.kind == "text":
                box.height = _estimate_text_height(box.shape.text_frame, box.shape, ctx, width)
            elif box.kind == "picture":
                box.pic_w, box.pic_h = _picture_box(box.shape, ctx, width)
                box.height = box.pic_h
            elif box.kind == "table":
                box.height = min(box.bottom - box.top, ctx.usable_height // 2)
            else:
                box.height = 20 * EMU_PER_PT

    def total_height(scale: float) -> int:
        h = 0 if title_boxes else HEADING_FALLBACK_HEIGHT
        for row in [title_boxes] + rows:
            if row:
                h += max((int(b.height * scale) if b.kind == "picture" else b.height) for b in row)
        return h

    available = int(ctx.usable_height * PAGE_FILL)
    scale = 1.0
    if total_height(1.0) > available:
        lo, hi = 0.25, 1.0
        for _ in range(12):
            mid = (lo + hi) / 2
            if total_height(mid) > available:
                hi = mid
            else:
                lo = mid
        scale = lo
    for row in rows:
        for box in row:
            if box.kind == "picture":
                box.pic_w = int(box.pic_w * scale)
                box.pic_h = int(box.pic_h * scale)
    ctx.gap_budget = max(0, available - total_height(scale))


# ---------- emitting boxes ----------


def _emit_text_frame(
    container, text_frame, shape, ctx: Ctx, first_paragraph=None, force_style=None, extra_before=0, shade=None
):
    """Writes the paragraphs of a text frame into `container` (document or table cell)."""
    scale = _font_scale(shape) if shape is not None else 1.0
    reuse = first_paragraph
    ctx.font_ref_color = _shape_style_color(shape, ctx, "a:fontRef") if shape is not None else None
    try:
        for paragraph in _trimmed_paragraphs(text_frame):
            _emit_paragraph(
                container,
                paragraph,
                shape,
                ctx,
                reuse=reuse,
                force_style=force_style,
                font_scale=scale,
                extra_before=extra_before,
                shade=shade,
            )
            reuse = None
            extra_before = 0
    finally:
        ctx.font_ref_color = None


def _emit_title(title_boxes: list[Box], ctx: Ctx, new_page: bool) -> None:
    h = ctx.doc.add_heading("", level=1)
    h.paragraph_format.page_break_before = new_page
    for box in title_boxes:
        if box.kind == "text":
            _emit_text_frame(
                ctx.doc,
                box.shape.text_frame,
                box.shape,
                ctx,
                first_paragraph=h,
                force_style="Heading 1",
                extra_before=_take_gap(ctx, box.top),
                shade=_shape_fill(box.shape, ctx),
            )
            _advance(ctx, box.top, box.bottom - box.top)
            return
    h.add_run(f"Slide {ctx.slide_no}")
    h.paragraph_format.space_before = Pt(0)
    ctx.report.note("Slides without a title got a 'Slide n' heading.")


def _tiny_paragraph(container) -> None:
    """Word needs a paragraph after a table; keep it as short as possible."""
    p = container.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    p.add_run("").font.size = Pt(1)
    _set_paragraph_mark_size(p, 1)


def _spacer(container, gap: int) -> None:
    if gap > 0:
        spacer = container.add_paragraph()
        spacer.paragraph_format.space_before = Emu(gap)
        spacer.paragraph_format.space_after = Pt(0)
        spacer.add_run("").font.size = Pt(1)


def _emit_table(table, container, ctx: Ctx, width: int) -> None:
    rows, cols = len(table.rows), len(table.columns)
    dt = container.add_table(rows, cols)
    dt.style = "Table Grid"
    dt.autofit = False
    total = sum(int(c.width or 0) for c in table.columns) or 1
    scale = min(1.0, width / total)
    for i, col in enumerate(table.columns):
        w = Emu(int((col.width or total / cols) * scale))
        dt.columns[i].width = w
        for cell in dt.columns[i].cells:
            cell.width = w
    merges = []
    for r in range(rows):
        for c in range(cols):
            cell = table.cell(r, c)
            if cell.is_spanned:
                continue
            dcell = dt.cell(r, c)
            _emit_text_frame(dcell, cell.text_frame, None, ctx, first_paragraph=dcell.paragraphs[0])
            if cell.is_merge_origin:
                merges.append((r, c, r + cell.span_height - 1, c + cell.span_width - 1))
    for r0, c0, r1, c1 in merges:
        dt.cell(r0, c0).merge(dt.cell(r1, c1))
    _tiny_paragraph(container)
    ctx.report.tables += 1


def _prepare_picture(shape):
    """Returns (bytes, width_emu, height_emu) as the picture appears on the slide.
    Word ignores EXIF orientation and knows nothing about PowerPoint's crop and rotation,
    so those are baked into the pixels with Pillow when needed."""
    image = shape.image
    data = image.blob
    w, h = int(shape.width or 0), int(shape.height or 0)
    rotation = float(shape.rotation or 0) % 360
    crops = [max(0.0, float(getattr(shape, f"crop_{side}", 0) or 0)) for side in ("left", "top", "right", "bottom")]
    needs_pixels = rotation or any(crops)
    with Image.open(BytesIO(data)) as im:
        if im.getexif().get(0x0112, 1) != 1:
            needs_pixels = True
        if not needs_pixels:
            return data, w, h
        im = ImageOps.exif_transpose(im)
        if any(crops):
            pw, ph = im.size
            box = (int(pw * crops[0]), int(ph * crops[1]), int(pw * (1 - crops[2])), int(ph * (1 - crops[3])))
            if box[2] > box[0] and box[3] > box[1]:
                im = im.crop(box)
        if rotation:
            im = im.rotate(-rotation, expand=True, fillcolor=(255, 255, 255, 0) if im.mode in ("RGBA", "LA") else None)
            if round(rotation) % 180 == 90:
                w, h = h, w
            elif round(rotation) % 90:
                # Arbitrary angle: keep the unrotated width and let the bounding box follow the rotated pixels.
                h = int(w * im.height / im.width) if im.width else h
        if im.mode not in ("RGB", "RGBA", "L", "LA", "P"):
            im = im.convert("RGBA")
        out = BytesIO()
        im.save(out, format="PNG")
        return out.getvalue(), w, h


def _emit_picture(box: Box, container, ctx: Ctx, first_paragraph=None, gap: int = 0) -> None:
    shape = box.shape
    ext = (shape.image.ext or "").lower()
    p = first_paragraph if first_paragraph is not None else container.add_paragraph()
    if ext not in RASTER_EXT:
        p.add_run(f"[Image skipped: {ext.upper() or 'vector'} pictures are not supported]").italic = True
        ctx.report.warn(
            f"Slide {ctx.slide_no}: a {ext.upper() or 'vector'} picture was skipped (EMF/WMF/SVG are not supported)."
        )
        return
    try:
        data, w, h = _prepare_picture(shape)
    except Exception as exc:  # unreadable pixels: fall back to the raw blob and original box
        log.warning("slide %d: could not preprocess picture %r (%s)", ctx.slide_no, shape.name, exc)
        data, w, h = shape.image.blob, int(shape.width or 0), int(shape.height or 0)
    width = box.pic_w or min(max(w, 1), box.col_width or ctx.usable_width)
    # Given only a width, Word keeps the image's own aspect ratio. PowerPoint lets a picture be stretched to its
    # frame, and _plan_slide budgeted the frame's height, so pass the height too.
    height = Emu(max(1, int(width * h / w))) if w > 0 and h > 0 else None
    try:
        p.add_run().add_picture(BytesIO(data), width=Emu(max(1, width)), height=height)
    except UnrecognizedImageError:
        ctx.report.warn(f"Slide {ctx.slide_no}: a picture could not be decoded and was skipped.")
        return
    p.paragraph_format.space_before = Emu(gap)
    p.paragraph_format.space_after = Pt(0)
    ctx.report.images += 1


def _emit_box(box: Box, container, ctx: Ctx, first_paragraph=None, gap: int = 0, cell=None) -> None:
    shape = box.shape
    if box.kind == "table":
        _spacer(container, gap)
        _emit_table(shape.table, container, ctx, box.col_width or ctx.usable_width)
    elif box.kind == "picture":
        _emit_picture(box, container, ctx, first_paragraph, gap)
    elif box.kind == "text":
        fill = _shape_fill(shape, ctx)
        if fill and cell is not None:
            _shade_cell(cell, fill)
            fill = None
        _emit_text_frame(
            container, shape.text_frame, shape, ctx, first_paragraph=first_paragraph, extra_before=gap, shade=fill
        )
    else:
        p = first_paragraph if first_paragraph is not None else container.add_paragraph()
        if getattr(shape, "has_chart", False) and shape.has_chart:
            title = ""
            try:
                if shape.chart.has_title:
                    title = shape.chart.chart_title.text_frame.text.strip()
            except Exception:
                title = ""
            p.add_run(f"[Chart: {title or shape.name}]").italic = True
            ctx.report.warn(f"Slide {ctx.slide_no}: chart '{title or shape.name}' was replaced by a placeholder.")
        else:
            p.add_run(f"[{shape.name}: SmartArt or embedded object not converted]").italic = True
            ctx.report.warn(
                f"Slide {ctx.slide_no}: '{shape.name}' (SmartArt or embedded object) was replaced by a placeholder."
            )


def _emit_row(row: list[Box], ctx: Ctx) -> None:
    top = min(b.top for b in row)
    bottom = max(b.bottom for b in row)
    gap = _take_gap(ctx, top)
    if len(row) == 1:
        _emit_box(row[0], ctx.doc, ctx, gap=gap)
    else:
        _spacer(ctx.doc, gap)
        layout = ctx.doc.add_table(1, len(row))
        layout.autofit = False
        for i, box in enumerate(row):
            cell = layout.cell(0, i)
            layout.columns[i].width = Emu(box.col_width)
            cell.width = Emu(box.col_width)
            # A shape that starts well below the row's top sat vertically centred beside its neighbour.
            if box.top - top > 0.2 * max(1, bottom - top):
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            try:
                _emit_box(box, cell, ctx, first_paragraph=cell.paragraphs[0], cell=cell)
            except Exception as exc:
                log.exception("slide %d shape %r failed", ctx.slide_no, box.shape.name)
                ctx.report.warn(
                    f"Slide {ctx.slide_no}: shape '{box.shape.name}' could not be converted ({exc.__class__.__name__})."
                )
        _tiny_paragraph(ctx.doc)
    _advance(ctx, top, bottom - top)


def _emit_notes(slide, ctx: Ctx) -> None:
    if not slide.has_notes_slide:
        return
    text = slide.notes_slide.notes_text_frame.text.strip()
    if not text:
        return
    ctx.doc.add_heading("Speaker notes", level=3)
    for line in text.splitlines():
        if line.strip():
            ctx.doc.add_paragraph().add_run(line.strip()).italic = True


# ---------- entry point ----------


def convert_pptx(src: Path, dst: Path, include_notes: bool = False) -> ConversionReport:
    try:
        prs = Presentation(str(src))
    except (PackageNotFoundError, KeyError, ValueError) as exc:
        raise CorruptFileError(f"Not a readable PPTX: {exc}") from exc
    report = ConversionReport()
    doc = Document()
    uw, uh = _setup_page(doc, prs, report)
    ctx = Ctx(doc=doc, report=report, usable_width=uw, usable_height=uh, prs=prs)
    slides = list(prs.slides)
    report.pages = len(slides)
    for idx, slide in enumerate(slides, start=1):
        ctx.slide_no = idx
        ctx.slide = slide
        master = slide.slide_layout.slide_master
        key = master.part.partname
        if key not in ctx.themes:
            ctx.themes[key] = Theme(prs, master)
        ctx.theme = ctx.themes[key]
        ctx.cursor = MARGIN
        titles, body = _ordered_shapes(slide)
        title_boxes = [b for b in _make_boxes(titles, ctx) if b.kind == "text"][:1]
        rows = _group_rows(_make_boxes(body, ctx))
        _plan_slide(title_boxes, rows, ctx)
        _emit_title(title_boxes, ctx, new_page=idx > 1)
        for row in rows:
            try:
                _emit_row(row, ctx)
            except Exception as exc:  # keep converting the rest of the deck
                names = ", ".join(b.shape.name for b in row)
                log.exception("slide %d shapes %s failed", idx, names)
                report.warn(f"Slide {idx}: shape '{names}' could not be converted ({exc.__class__.__name__}).")
        if include_notes:
            _emit_notes(slide, ctx)
    report.note("Animations, transitions and slide backgrounds are not carried over.")
    doc.save(str(dst))
    return report
